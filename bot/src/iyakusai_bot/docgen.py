# -*- coding: utf-8 -*-
"""ファイル出力(F4)。Markdown テキスト → docx / md / txt のバイト列。

python-docx は純 Python なので RPi / Windows どちらでも動く。
Word COM 等の OS 依存経路は使わない(マルチホスト要件 §4.1)。
対応 Markdown は見出し(#〜###)・箇条書き(- / 1.)・太字(**)・表(|区切り)・段落。
"""

from __future__ import annotations

import io
import re
from datetime import datetime
from zoneinfo import ZoneInfo

FORMATS = ("docx", "md", "txt")
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_FNAME_BAD = re.compile(r'[\\/:*?"<>|\s]+')


def sanitize_filename(title: str, fmt: str) -> str:
    stem = _FNAME_BAD.sub("_", title).strip("_")[:40] or "document"
    stamp = datetime.now(ZoneInfo("Asia/Tokyo")).strftime("%Y%m%d")
    return f"{stem}_{stamp}.{fmt}"


def to_file(fmt: str, title: str, body_md: str) -> tuple[str, bytes]:
    """(ファイル名, バイト列) を返す。fmt は FORMATS のいずれか。"""
    if fmt == "md":
        return sanitize_filename(title, "md"), body_md.encode("utf-8")
    if fmt == "txt":
        plain = _BOLD_RE.sub(r"\1", body_md)
        plain = re.sub(r"^#{1,4}\s*", "", plain, flags=re.M)
        return sanitize_filename(title, "txt"), plain.encode("utf-8")
    if fmt == "docx":
        return sanitize_filename(title, "docx"), _build_docx(title, body_md)
    raise ValueError(f"未対応の形式: {fmt}")


def _build_docx(title: str, body_md: str) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    # 和文フォント(委員会の既存資料に合わせて游ゴシック)
    style = doc.styles["Normal"]
    style.font.name = "游ゴシック"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "游ゴシック")

    lines = body_md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line.strip() or re.fullmatch(r"-{3,}|\*{3,}", line.strip()):
            continue  # 空行・水平線はスキップ
        if line.lstrip().startswith("|"):
            # 表ブロックをまとめて取り込む
            rows = [line]
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(lines[i].rstrip())
                i += 1
            _add_table(doc, rows)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            level = min(len(m.group(1)), 3)
            doc.add_heading(_BOLD_RE.sub(r"\1", m.group(2)), level=level)
            continue
        m = re.match(r"^[-*]\s+(.*)", line)
        if m:
            _add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
            continue
        m = re.match(r"^\d+[.)]\s+(.*)", line)
        if m:
            _add_runs(doc.add_paragraph(style="List Number"), m.group(1))
            continue
        _add_runs(doc.add_paragraph(), line)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_table(doc, rows: list[str]) -> None:
    """Markdown 表(| セル | セル |)を Word の表に変換する。"""
    parsed = []
    for r in rows:
        cells = [c.strip() for c in r.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            continue  # 区切り行(---)
        parsed.append(cells)
    if not parsed:
        return
    ncols = max(len(r) for r in parsed)
    table = doc.add_table(rows=len(parsed), cols=ncols)
    table.style = "Table Grid"
    for ri, cells in enumerate(parsed):
        for ci in range(ncols):
            text = cells[ci] if ci < len(cells) else ""
            cell_par = table.cell(ri, ci).paragraphs[0]
            _add_runs(cell_par, text)
            if ri == 0:  # 1行目は見出し扱いで太字
                for run in cell_par.runs:
                    run.bold = True


def _add_runs(paragraph, text: str) -> None:
    """**太字** を run に分解して追加する。"""
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        paragraph.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])
